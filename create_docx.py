import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def md_to_docx(md_path, docx_path):
    doc = docx.Document()
    
    # Page Margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()
        
    in_table = False
    table_data = []
    
    for line in lines:
        raw_line = line.rstrip('\n')
        stripped = raw_line.strip()
        
        if stripped.startswith('|') and stripped.endswith('|'):
            # Table row
            parts = [p.strip() for p in stripped.split('|')[1:-1]]
            if all(set(p).issubset(set('-: ')) for p in parts if p):
                continue # Header separator
            table_data.append(parts)
            in_table = True
            continue
        else:
            if in_table and table_data:
                # Flush table
                table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                table.style = 'Table Grid'
                for r_idx, row in enumerate(table_data):
                    for c_idx, cell_value in enumerate(row):
                        if c_idx < len(table.columns):
                            cell = table.cell(r_idx, c_idx)
                            cell.text = cell_value
                            # If header row
                            if r_idx == 0:
                                shading = parse_xml(r'<w:shd {} w:fill="1F4E79"/>'.format(nsdecls('w')))
                                cell._tc.get_or_add_tcPr().append(shading)
                                for p in cell.paragraphs:
                                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                                    for r in p.runs:
                                        r.font.bold = True
                                        r.font.color.rgb = RGBColor(255, 255, 255)
                doc.add_paragraph() # spacer
                table_data = []
                in_table = False

        if not stripped:
            continue
            
        if stripped.startswith('# '):
            p = doc.add_paragraph()
            r = p.add_run(stripped[2:])
            r.font.size = Pt(20)
            r.font.bold = True
            r.font.color.rgb = RGBColor(31, 78, 121)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif stripped.startswith('## '):
            p = doc.add_paragraph()
            r = p.add_run(stripped[3:])
            r.font.size = Pt(14)
            r.font.bold = True
            r.font.color.rgb = RGBColor(46, 117, 182)
        elif stripped.startswith('### '):
            p = doc.add_paragraph()
            r = p.add_run(stripped[4:])
            r.font.size = Pt(12)
            r.font.bold = True
            r.font.color.rgb = RGBColor(89, 89, 89)
        elif stripped.startswith('- ') or stripped.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            text = stripped[2:]
            # parse bold inside
            parse_runs(p, text)
        elif stripped.startswith('1. ') or stripped.startswith('2. ') or stripped.startswith('3. ') or stripped.startswith('4. ') or stripped.startswith('5. '):
            p = doc.add_paragraph(style='List Number')
            text = stripped.split('. ', 1)[1]
            parse_runs(p, text)
        else:
            p = doc.add_paragraph()
            parse_runs(p, stripped)

    if in_table and table_data:
        table = doc.add_table(rows=len(table_data), cols=len(table_data[0]))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.style = 'Table Grid'
        for r_idx, row in enumerate(table_data):
            for c_idx, cell_value in enumerate(row):
                if c_idx < len(table.columns):
                    cell = table.cell(r_idx, c_idx)
                    cell.text = cell_value
                    if r_idx == 0:
                        shading = parse_xml(r'<w:shd {} w:fill="1F4E79"/>'.format(nsdecls('w')))
                        cell._tc.get_or_add_tcPr().append(shading)
                        for p in cell.paragraphs:
                            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            for r in p.runs:
                                r.font.bold = True
                                r.font.color.rgb = RGBColor(255, 255, 255)
                                
    doc.save(docx_path)

def parse_runs(p, text):
    parts = text.split('**')
    for idx, part in enumerate(parts):
        if not part:
            continue
        run = p.add_run(part)
        if idx % 2 == 1:
            run.font.bold = True

if __name__ == '__main__':
    md_to_docx(sys.argv[1], sys.argv[2])
